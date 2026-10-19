"""Downloadable reports: the same summary and tables as a page section, as an
Excel workbook (one sheet per table, numbers kept as numbers so they can be
filtered and summed) and as a Word document (headings, the summary, then
each table) for printing or sending on.

A page describes its report once — Report(title, subtitle, summary, tables)
— and asks for either format, so the two never drift apart.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BRAND = "1C5CAB"   # theme.SEQUENTIAL_BLUE step: dark enough for white header text
_MONEY = re.compile(r"revenue|sales|target|value|amount|kes|takings|per day|per attendant|avg sale",
                    re.I)
_PCT = re.compile(r"%")


@dataclass
class Table:
    name: str
    data: pd.DataFrame
    note: str | None = None
    # Word pages are narrow: a wide table is cut to these columns there; the
    # workbook always gets every column.
    doc_columns: list[str] | None = None
    # Long row-level tables (every POS session) belong in the workbook only.
    in_doc: bool = True


@dataclass
class Report:
    title: str
    subtitle: str
    summary: list[tuple[str, str]] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def _kind(column: str, series: pd.Series) -> str:
    if not pd.api.types.is_numeric_dtype(series) or pd.api.types.is_bool_dtype(series):
        return "text"
    if _PCT.search(column):
        return "pct"
    if _MONEY.search(column):
        return "money"
    if (series.dropna() % 1 == 0).all():
        return "int"
    return "dec"


def _cell_text(value, kind: str) -> str:
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return "—"
    if kind == "money":
        return f"{value:,.0f}"
    if kind == "pct":
        return f"{value:.1f}%"
    if kind == "int":
        return f"{value:,.0f}"
    if kind == "dec":
        return f"{value:,.2f}"
    return str(value)


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "", name)[:31] or "Sheet"
    out, n = base, 2
    while out in used:
        out = f"{base[:28]} {n}"
        n += 1
    used.add(out)
    return out


def to_xlsx(report: Report) -> bytes:
    buf = io.BytesIO()
    head_fill = PatternFill("solid", fgColor=BRAND)
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D0D0D0")
    used: set[str] = set()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        summary = pd.DataFrame(report.summary, columns=["Measure", "Value"])
        summary.to_excel(xl, sheet_name=_sheet_name("Summary", used), index=False, startrow=3)
        ws = xl.sheets["Summary"]
        ws["A1"], ws["A2"] = report.title, report.subtitle
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"].font = Font(italic=True, color="666666")
        for c in ws[4]:
            c.fill, c.font = head_fill, head_font
        row = 6 + len(summary)
        for note in report.notes:
            ws.cell(row=row, column=1, value=note).alignment = Alignment(wrap_text=False)
            row += 1
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 44

        for table in report.tables:
            name = _sheet_name(table.name, used)
            start = 2 if table.note else 0
            table.data.to_excel(xl, sheet_name=name, index=False, startrow=start)
            ws = xl.sheets[name]
            if table.note:
                ws["A1"] = table.note
                ws["A1"].font = Font(italic=True, color="666666")
            header = start + 1
            for c in ws[header]:
                c.fill, c.font = head_fill, head_font
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(table.data.columns, start=1):
                kind = _kind(col, table.data[col])
                fmt = {"money": "#,##0", "pct": '0.0"%"', "int": "#,##0", "dec": "#,##0.00"}.get(kind)
                letter = get_column_letter(i)
                for cell in ws[letter][header:]:
                    cell.border = Border(bottom=thin)
                    if fmt:
                        cell.number_format = fmt
                longest = max([len(str(col))] + [len(_cell_text(v, kind)) for v in table.data[col].head(500)])
                ws.column_dimensions[letter].width = min(max(10, longest + 2), 48)
            ws.freeze_panes = ws.cell(row=header + 1, column=2)
            ws.auto_filter.ref = (f"A{header}:{get_column_letter(len(table.data.columns))}"
                                  f"{header + len(table.data)}")
    return buf.getvalue()


def _shade(cell, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    props.append(shd)


def to_docx(report: Report) -> bytes:
    doc = Document()
    section = doc.sections[0]
    # Landscape: the ranking and score tables are wide.
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.6))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading(report.title, level=0)
    sub = doc.add_paragraph(report.subtitle)
    sub.runs[0].italic = True
    sub.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    if report.summary:
        doc.add_heading("Summary", level=1)
        t = doc.add_table(rows=0, cols=2)
        t.style = "Light List Accent 1"
        for label, value in report.summary:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = label, value
            cells[0].paragraphs[0].runs[0].bold = True
    for note in report.notes:
        p = doc.add_paragraph(note)
        p.runs[0].font.size = Pt(9)

    left_out = [t.name for t in report.tables if not t.in_doc]
    for table in (t for t in report.tables if t.in_doc):
        doc.add_heading(table.name, level=1)
        if table.note:
            p = doc.add_paragraph(table.note)
            p.runs[0].font.size = Pt(9)
            p.runs[0].italic = True
        data = table.data[table.doc_columns] if table.doc_columns else table.data
        if data.empty:
            doc.add_paragraph("Nothing to show for this selection.")
            continue
        kinds = {c: _kind(c, data[c]) for c in data.columns}
        t = doc.add_table(rows=1, cols=len(data.columns))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, col in zip(t.rows[0].cells, data.columns):
            cell.text = str(col)
            run = cell.paragraphs[0].runs[0]
            run.bold, run.font.size = True, Pt(8.5)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell, BRAND)
        for values in data.itertuples(index=False):
            cells = t.add_row().cells
            for cell, col, value in zip(cells, data.columns, values):
                cell.text = _cell_text(value, kinds[col])
                para = cell.paragraphs[0]
                para.runs[0].font.size = Pt(8.5)
                if kinds[col] != "text":
                    para.alignment = 2  # right
        if table.doc_columns and len(table.doc_columns) < len(table.data.columns):
            p = doc.add_paragraph(f"Key columns only — the Excel report has all "
                                  f"{len(table.data.columns)}.")
            p.runs[0].font.size = Pt(8)

    if left_out:
        p = doc.add_paragraph("In the Excel report only (too long to print): " + ", ".join(left_out) + ".")
        p.runs[0].font.size = Pt(8)
    foot = doc.add_paragraph(f"Generated {datetime.now():%d %b %Y %H:%M} · Denri Sales Operations Dashboard")
    foot.runs[0].font.size = Pt(8)
    foot.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
